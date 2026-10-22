#!/usr/bin/env python3
"""Read the two owner-provided workbooks into an ignored, private local batch.

Only aggregate counts go to stdout. Source workbooks are never modified. Names
are deliberately neither printed nor used as person identity keys.
"""
import argparse
import collections
import hashlib
import json
import os
import re
import secrets
import time
from pathlib import Path


def clean(value):
    return "" if value is None else str(value).strip()


def stable_id(prefix, key):
    return prefix + hashlib.sha256(key.encode()).hexdigest()[:16]


def suggested_subjects(job):
    subjects = [subject for subject in ["数学", "物理", "化学", "英语", "经济", "生物", "语文", "体育", "历史", "音乐", "粤语"] if subject in job]
    if "艺术" in job or "美术" in job:
        subjects.append("艺术")
    return subjects


def parse_staff(sheets, batch):
    people = []
    for sheet_name, rows in sheets:
        header_index = next((i for i, row in enumerate(rows[:10]) if "姓名" in [clean(v) for v in row] and "部门" in [clean(v) for v in row]), None)
        if header_index is None:
            raise ValueError("Personnel sheet has no supported header")
        headers = [clean(v) for v in rows[header_index]]
        name_col, dept_col = headers.index("姓名"), headers.index("部门")
        job_col = next((headers.index(k) for k in ["职位", "职务"] if k in headers), None)
        serial_col = headers.index("序号") if "序号" in headers else None
        for row_index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
            get = lambda i: clean(row[i]) if i is not None and i < len(row) else ""
            name = get(name_col)
            if not name:
                continue
            department = get(dept_col) or sheet_name
            job = get(job_col)
            # Source serials need not be unique (e.g. job rows share a serial).
            # Keep row identity instead of merging either serials or names.
            key = f"{sheet_name}|row-{row_index}|{get(serial_col)}"
            people.append({
                "id": stable_id("local-staff-", key), "name": name, "employeeNo": "",
                "department": department, "jobTitle": job, "gender": "未透露",
                "phone": "", "email": "", "status": "active", "systemRoles": [],
                "duties": [], "accountStatus": "none", "history": [],
                "sourceBatch": batch, "sourceRefs": [f"教职工列表 / {sheet_name}!A{row_index}:D{row_index}"],
                "departmentAffiliations": [department], "suggestedResearchSubjects": suggested_subjects(job),
            })
    if len({p["id"] for p in people}) != len(people):
        raise ValueError("Duplicate personnel source key; no batch written")
    frequencies = collections.Counter(p["name"] for p in people)
    for person in people:
        person["duplicateName"] = frequencies[person["name"]] > 1
    return people


def place_kind(usage):
    if not usage or "空置" in usage:
        return "other"
    if "实验" in usage:
        return "lab"
    if "办公室" in usage or usage in ["艺术组", "数学组", "英语组", "物理组", "化学组", "经济组", "通识组"]:
        return "office"
    if "会议" in usage or "报告厅" in usage or "面谈" in usage:
        return "meeting"
    if "班" in usage or "教室" in usage or re.match(r"Pre\s", usage, re.I):
        return "classroom"
    return "function"


def parse_places(sheet_name, cells, merged_ranges, batch):
    """cells uses (row, col) keys; merged ranges are inclusive row/col bounds."""
    def merged_value(row, col):
        if cells.get((row, col)) is not None:
            return clean(cells[(row, col)])
        for r1, c1, r2, c2 in merged_ranges:
            if r1 <= row <= r2 and c1 <= col <= c2:
                return clean(cells.get((r1, c1)))
        return ""

    places = {}
    ignored_markers = 0
    for (row, col), value in sorted(cells.items()):
        if row < 3 or col not in [4, 6] or not clean(value):
            continue
        code = merged_value(row, col - 1)
        if not re.fullmatch(r"[A-Z]\d{3}", code):
            raise ValueError("A usage has no supported room code; no batch written")
        area = merged_value(row, 2)
        if not area:
            raise ValueError("Room has no explicit area; no batch written")
        key = area + "|" + code
        usage = clean(value)
        ref = f"场所分布图 / {sheet_name}!{'C' if col == 4 else 'E'}{row}:{'D' if col == 4 else 'F'}{row}"
        if key not in places:
            places[key] = {"id": stable_id("local-place-", key), "name": code, "roomCode": code,
                           "building": area, "area": area, "floor": "", "status": "active",
                           "order": len(places) * 10 + 1000, "usage": usage, "kind": place_kind(usage),
                           "sourceRef": ref, "sourceBatch": batch, "note": ""}
        else:
            place = places[key]
            place["usage"] += " / " + usage
            place["sourceRef"] += "；" + ref
            place["note"] = "原表同一门牌包含多项用途，保留为一个场所，待核对是否需要分室。"
    for (row, col), value in cells.items():
        if row >= 3 and col in [3, 5] and clean(value) and not re.fullmatch(r"[A-Z]\d{3}", clean(value)):
            ignored_markers += 1
    return list(places.values()), ignored_markers


def assemble(staff_sheets, place_sheet, cells, merges, batch):
    staff = parse_staff(staff_sheets, batch)
    places, markers = parse_places(place_sheet, cells, merges, batch)
    departments = collections.Counter(p["department"] for p in staff)
    jobs = collections.Counter(p["jobTitle"] for p in staff if p["jobTitle"])
    subjects = collections.Counter(subject for p in staff for subject in p["suggestedResearchSubjects"])
    return {
        "schema": 1, "batchId": batch,
        "staff": staff,
        "departments": [{"id": stable_id("local-dept-", name), "name": name, "status": "active", "order": i * 10 + 1000, "referencedBy": count, "sourceBatch": batch, "note": "来自本机人员导入"} for i, (name, count) in enumerate(departments.items())],
        "jobs": [{"id": stable_id("local-job-", name), "name": name, "status": "active", "order": i * 10 + 1000, "referencedBy": count, "sourceBatch": batch} for i, (name, count) in enumerate(jobs.items())],
        "places": places,
        "summary": {"staffRows": len(staff), "departments": dict(departments), "jobs": len(jobs), "places": len(places), "areas": dict(collections.Counter(p["area"] for p in places)), "placeKinds": dict(collections.Counter(p["kind"] for p in places)), "suggestedSubjects": dict(subjects), "duplicateNameRows": sum(p["duplicateName"] for p in staff), "multiUsePlaces": sum(bool(p["note"]) for p in places), "ignoredLayoutMarkers": markers},
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("staff_workbook", type=Path)
    parser.add_argument("places_workbook", type=Path)
    args = parser.parse_args()
    import openpyxl
    source_hash = hashlib.sha256(args.staff_workbook.read_bytes() + args.places_workbook.read_bytes()).hexdigest()
    batch = "school-files-" + source_hash[:16]
    staff_book = openpyxl.load_workbook(args.staff_workbook, data_only=True)
    place_book = openpyxl.load_workbook(args.places_workbook, data_only=True)
    sheet = place_book.worksheets[0]
    bundle = assemble([(s.title, list(s.values)) for s in staff_book], sheet.title,
                      {(c.row, c.column): c.value for row in sheet for c in row if c.value is not None},
                      [(r.min_row, r.min_col, r.max_row, r.max_col) for r in sheet.merged_cells.ranges], batch)
    destination = Path(__file__).resolve().parents[1] / ".local-artifacts/private-school-import"
    destination.mkdir(parents=True, exist_ok=True)
    os.chmod(destination, 0o700)
    target = destination / "current.json"
    # Existing batches are backed up privately before the explicit new import.
    if target.exists() and target.read_text() != json.dumps(bundle, ensure_ascii=False, indent=2):
        previous = target.read_bytes()
        backup = destination / ("previous-" + hashlib.sha256(previous).hexdigest()[:12] + ".json")
        backup.write_bytes(previous)
        os.chmod(backup, 0o600)
    target.write_text(json.dumps(bundle, ensure_ascii=False, indent=2))
    os.chmod(target, 0o600)
    # Only the owner launch fragment receives the bearer capability. The server
    # checks its digest; no token is placed in a URL query, output, or source.
    token = secrets.token_hex(32)
    capability = destination / "capability.json"
    capability.write_text(json.dumps({"digest": hashlib.sha256(token.encode()).hexdigest(), "batchId": batch, "expiresAt": int(time.time()) + 3 * 86400}))
    os.chmod(capability, 0o600)
    launch = destination / "launch-token"
    launch.write_text(token)
    os.chmod(launch, 0o600)
    print(json.dumps({"batchId": batch, **bundle["summary"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
